"""Damping force: regra lida da planilha 'DAMPING FORCE - REGRA' (tolerâncias, velocidades e a fórmula min/máx)."""
import re, datetime, io
import openpyxl
import database as db

NOTE = ("Na planilha original a verificação do limite superior usa F>=G na 1ª linha de cada tabela e F>=-G nas demais; com -G o limite superior nunca é testado. "
        "Aqui é aplicada a mesma condição em todas as linhas: mín <= ENCONTRADO <= máx.")

def parse_rules(path):
    ws = openpyxl.load_workbook(path).worksheets[0]; out, params = [], {}
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(r, 3).value or "").strip().upper() != "VM/S": continue
        m = re.match(r"\s*([A-Za-z]+)", str(ws.cell(r, 5).value or ""))
        if not m: raise ValueError(f"Linha {r}: cabeçalho de força não reconhecido.")
        name, seq, rr = m.group(1), 0, r + 1
        while isinstance(ws.cell(rr, 3).value, (int, float)) and isinstance(ws.cell(rr, 2).value, (int, float)):
            out.append((name, seq, float(ws.cell(rr, 3).value), float(ws.cell(rr, 2).value))); seq += 1; rr += 1
            f = str(ws.cell(rr - 1, 5).value or "")
            g = re.search(r"IF\(D\d+<(\d+(?:\.\d+)?),\(D\d+\*\(1-B\d+\)\)\+(\d+(?:\.\d+)?),", f)
            if g and "threshold" not in params: params = {"threshold": float(g.group(1)), "offset": float(g.group(2))}
    if not out or not params: raise ValueError("Não encontrei as tabelas (coluna VM/S) ou a fórmula IF(D<limite, …+ajuste) na planilha.")
    return out, params

def load_rules(path, filename):
    rules, params = parse_rules(path); c = db.conn()
    try:
        r = c.execute("SELECT id FROM sources WHERE name=?", (filename,)).fetchone()
        src = r[0] if r else c.execute("INSERT INTO sources(name,category) VALUES(?,?)", (filename, "Regra de damping force")).lastrowid
        c.execute("DELETE FROM damping_rules"); c.executemany("INSERT INTO damping_rules(table_name,seq,velocity,tolerance,source_id) VALUES(?,?,?,?,?)", [(n, s, v, t, src) for n, s, v, t in rules])
        for k, v in params.items(): c.execute("INSERT INTO damping_params(key,value,note) VALUES(?,?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (k, v, filename))
        c.execute("INSERT INTO import_batches(source_id,filename,category,rows_new,rows_read,status) VALUES(?,?,?,?,?,'ok')", (src, filename, "Regra de damping force", len(rules), len(rules)))
        c.commit(); return len(rules), params
    finally: c.close()

def rules():
    c = db.conn()
    try:
        rs = [dict(r) for r in c.execute("SELECT r.table_name,r.seq,r.velocity,r.tolerance,s.name source FROM damping_rules r LEFT JOIN sources s ON s.id=r.source_id ORDER BY r.table_name,r.seq")]
        p = {r["key"]: r["value"] for r in c.execute("SELECT key,value FROM damping_params")}
    finally: c.close()
    tables = {}
    for r in rs: tables.setdefault(r["table_name"], []).append(r)
    return {"tables": tables, "threshold": p.get("threshold"), "offset": p.get("offset"), "note": NOTE, "source": rs[0]["source"] if rs else None}

def calc(drawing, tol, found, thr, off):
    """E = D*(1-tol) (+off se D<limite); F = D*(1+tol) (-off se D<limite); aprovado se E <= encontrado <= F."""
    if drawing is None: return None, None, None
    lo, hi = drawing * (1 - tol), drawing * (1 + tol)
    if drawing < thr: lo, hi = lo + off, hi - off
    if lo > hi: return lo, hi, "FAIXA INVÁLIDA"
    if found is None: return lo, hi, "INCOMPLETO"
    return lo, hi, "APPROVED" if lo <= found <= hi else "DISAPPROVED"

def overall(points):
    res = [p["result"] for p in points if p["result"]]
    if not res: return "INCOMPLETO"
    if "DISAPPROVED" in res: return "DISAPPROVED"
    if any(r in ("INCOMPLETO", "FAIXA INVÁLIDA") for r in res) or len(res) < len(points): return "INCOMPLETO"
    return "APPROVED"

def _fnum(v):
    if v in (None, ""): return None
    return float(str(v).replace(",", "."))

def get_test(tid):
    c = db.conn()
    try:
        t = c.execute("SELECT * FROM damping_tests WHERE id=?", (tid,)).fetchone()
        if not t: return None
        t = dict(t); pts = [dict(r) for r in c.execute("SELECT table_name,velocity,tolerance,drawing,found FROM damping_points WHERE test_id=? ORDER BY table_name,velocity", (tid,))]
        pr = {r["key"]: r["value"] for r in c.execute("SELECT key,value FROM damping_params")}
    finally: c.close()
    for p in pts: p["min"], p["max"], p["result"] = calc(p["drawing"], p["tolerance"], p["found"], pr["threshold"], pr["offset"])
    t["points"] = pts; t["overall"] = overall([p for p in pts if p["drawing"] is not None]); return t

def save_test(d):
    ru = rules()
    if not ru["tables"]: raise ValueError("Regra não carregada. Carregue a planilha DAMPING FORCE - REGRA.")
    tol = {(n, round(r["velocity"], 4)): r["tolerance"] for n, rs in ru["tables"].items() for r in rs}
    code = (d.get("perfect_code") or "").strip() or None; c = db.conn()
    try:
        pr = c.execute("SELECT id FROM products WHERE perfect_code=?", (code,)).fetchone() if code else None
        vals = (pr[0] if pr else None, code, d.get("label"), d.get("test_date") or datetime.date.today().isoformat(), d.get("notes"))
        if d.get("id"):
            c.execute("UPDATE damping_tests SET product_id=?,perfect_code=?,label=?,test_date=?,notes=? WHERE id=?", vals + (d["id"],)); tid = d["id"]; c.execute("DELETE FROM damping_points WHERE test_id=?", (tid,))
        else: tid = c.execute("INSERT INTO damping_tests(product_id,perfect_code,label,test_date,notes) VALUES(?,?,?,?,?)", vals).lastrowid
        for p in d.get("points") or []:
            key = (p["table"], round(float(p["velocity"]), 4))
            if key not in tol: raise ValueError(f"Velocidade {p['velocity']} da tabela {p['table']} não existe na regra.")
            c.execute("INSERT INTO damping_points(test_id,table_name,velocity,tolerance,drawing,found) VALUES(?,?,?,?,?,?)", (tid, p["table"], key[1], tol[key], _fnum(p.get("drawing")), _fnum(p.get("found"))))
        c.commit(); return tid
    finally: c.close()

def tests(q=None):
    c = db.conn()
    try: ids = [r[0] for r in c.execute("SELECT id FROM damping_tests WHERE (?1 IS NULL OR perfect_code LIKE ?2 OR label LIKE ?2) ORDER BY id DESC", (q or None, f"%{q}%"))]
    finally: c.close()
    return [{k: t[k] for k in ("id", "perfect_code", "label", "test_date", "overall", "product_id")} | {"points_filled": sum(1 for p in t["points"] if p["found"] is not None)} for t in map(get_test, ids)]

def prefill(code):
    c = db.conn()
    try: r = c.execute("SELECT id FROM damping_tests WHERE perfect_code=? ORDER BY id DESC LIMIT 1", (code,)).fetchone()
    finally: c.close()
    return get_test(r[0])["points"] if r else []

def test_workbook(tid):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    t = get_test(tid); wb = Workbook(); ws = wb.active; ws.title = "Damping force"
    ws["B1"] = "DAMPING FORCE CALCULATION"; ws["B1"].font = Font(bold=True, size=13)
    ws["B2"] = f"Produto: {t['perfect_code'] or '—'}   |   {t['label'] or ''}   |   Data: {t['test_date']}   |   Resultado geral: {t['overall']}"
    r = 4
    for name in sorted({p["table_name"] for p in t["points"]}):
        for i, h in enumerate(["PERFECT STANDARD", "VM/S", "DRAWING", f"{name} (N) min", f"{name} (N) Max", "FOUND", "RESULT"]):
            c = ws.cell(r, 2 + i, h); c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="243242"); c.alignment = Alignment(horizontal="center", wrap_text=True)
        for p in [x for x in t["points"] if x["table_name"] == name]:
            r += 1
            for i, v in enumerate([p["tolerance"], p["velocity"], p["drawing"], None if p["min"] is None else round(p["min"], 2), None if p["max"] is None else round(p["max"], 2), p["found"], p["result"]]):
                c = ws.cell(r, 2 + i, v)
                if i == 0: c.number_format = "0%"
                if i == 6: c.font = Font(bold=True, color={"APPROVED": "1E8E3E", "DISAPPROVED": "C5221F"}.get(v, "B06000"))
        r += 3
    ws.cell(r, 2, "Regra: se DRAWING < limite: mín = D×(1−tol)+ajuste e máx = D×(1+tol)−ajuste; senão sem ajuste. Aprovado se mín ≤ FOUND ≤ máx.").font = Font(italic=True)
    ws.cell(r + 1, 2, NOTE).font = Font(italic=True)
    for col, w in zip("BCDEFGH", [18, 10, 12, 14, 14, 12, 16]): ws.column_dimensions[col].width = w
    bio = io.BytesIO(); wb.save(bio); bio.seek(0); return bio
