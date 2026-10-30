"""Sugestões de equivalência: SIMILARIDADE de dados, nunca prova de intercambialidade. Nada é aprovado automaticamente."""
import json
import database as db

def _score(a, b, D, C, A, pairs):
    crit, diff, miss, sc, critical, srcs = [], [], [], 0, False, set()
    shared = C.get(a, {}).keys() & C.get(b, {}).keys()
    if shared: sc += 30; crit.append("Código em comum: " + ", ".join(sorted(shared)[:3]))
    aa, ab = A.get(a, []), A.get(b, [])
    if aa and ab:
        ov = [x for x in aa for y in ab if x["m"] and x["m"] == y["m"] and (x["f"] or 0) <= (y["t"] or 9999) and (y["f"] or 0) <= (x["t"] or 9999)]
        if ov: sc += 20; crit.append("Aplicação em comum: " + ov[0]["m"])
        else: diff.append("Nenhuma aplicação em comum")
        ea, eb = {x["ax"] for x in aa if x["ax"]}, {x["ax"] for x in ab if x["ax"]}
        if ea and eb:
            if ea & eb: sc += 5; crit.append("Mesmo eixo")
            else: diff.append("Eixo diferente"); critical = True
        else: miss.append("Eixo")
        sa, sb = {x["sd"] for x in aa if x["sd"]}, {x["sd"] for x in ab if x["sd"]}
        if sa and sb:
            if sa & sb or "ambos" in sa | sb: sc += 5; crit.append("Mesmo lado")
            else: diff.append("Lados diferentes"); critical = True
        else: miss.append("Lado")
    else: miss.append("Aplicações (um dos produtos não tem)")
    if (a, b) in pairs: diff.append("São lados opostos (direito × esquerdo) do mesmo conjunto"); critical = True
    da, dbb = D.get(a, {}), D.get(b, {})
    for key, lab in (("lmax_mm", "Lmax"), ("lmin_mm", "Lmin")):
        x, y = da.get(key), dbb.get(key)
        if x is None or y is None: miss.append(lab); continue
        dl = abs(x - y)
        if dl <= 2: sc += 10; crit.append(f"{lab} igual (Δ {dl:g} mm)")
        elif dl <= 5: sc += 5; crit.append(f"{lab} próximo (Δ {dl:g} mm)"); diff.append(f"{lab} difere {dl:g} mm")
        else: diff.append(f"{lab} difere {dl:g} mm"); critical = True
    sa, sb = da.get("system_type"), dbb.get("system_type")
    if sa and sb:
        if sa == sb: sc += 10; crit.append(f"Mesmo sistema ({sa})")
        else: diff.append(f"Sistema diferente ({sa} × {sb})"); critical = True
    else: miss.append("Sistema óleo/gás")
    fa, fb = da.get("family"), dbb.get("family")
    if fa and fb:
        if fa == fb: sc += 5; crit.append("Mesma família")
        else: diff.append(f"Família diferente ({fa} × {fb})")
    miss.append("Diâmetro da haste, tubo e fixações (não cadastrados)")
    conf = "Baixa" if critical else "Alta" if sc >= 75 else "Média" if sc >= 55 else "Baixa"
    return sc, conf, crit, diff, miss

def generate(min_score=50):
    c = db.conn()
    try:
        D = {r["id"]: dict(r) for r in c.execute("SELECT p.id,p.system_type,p.family,d.lmin_mm,d.lmax_mm FROM products p LEFT JOIN dimensions d ON d.product_id=p.id")}
        C, A = {}, {}
        for r in c.execute("SELECT product_id,normalized_code FROM codes WHERE normalized_code IS NOT NULL AND code_type IN ('oem','concorrente','fornecedor')"): C.setdefault(r[0], {})[r[1]] = 1
        for r in c.execute("SELECT product_id,lower(trim(model)),year_from,year_to,axle,side FROM applications"): A.setdefault(r[0], []).append({"m": r[1] or None, "f": r[2], "t": r[3], "ax": r[4], "sd": r[5]})
        pairs = {(r[0], r[1]) for r in c.execute("SELECT right_product_id,left_product_id FROM side_pairs")}; pairs |= {(b, a) for a, b in pairs}
        cand = {tuple(sorted(r)) for r in c.execute("SELECT a.product_id,b.product_id FROM codes a JOIN codes b ON a.normalized_code=b.normalized_code AND a.product_id<b.product_id WHERE a.normalized_code IS NOT NULL")}
        cand |= {tuple(r) for r in c.execute("""SELECT DISTINCT a.product_id,b.product_id FROM applications a JOIN applications b ON lower(trim(a.model))=lower(trim(b.model)) AND a.product_id<b.product_id
              AND coalesce(a.axle,'')=coalesce(b.axle,'') AND coalesce(a.model,'')<>'' AND coalesce(a.year_from,0)<=coalesce(b.year_to,9999) AND coalesce(b.year_from,0)<=coalesce(a.year_to,9999) LIMIT 30000""")}
        n = 0; kept = set()
        for a, b in cand:
            sc, conf, crit, diff, miss = _score(a, b, D, C, A, pairs)
            if sc < min_score: continue
            c.execute("""INSERT INTO equivalence_suggestions(product_a,product_b,score,confidence,criteria,differences,missing,sources,status,updated_at) VALUES(?,?,?,?,?,?,?,?,'pendente',CURRENT_TIMESTAMP)
                ON CONFLICT(product_a,product_b) DO UPDATE SET score=excluded.score,confidence=excluded.confidence,criteria=excluded.criteria,differences=excluded.differences,missing=excluded.missing,sources=excluded.sources,updated_at=CURRENT_TIMESTAMP""",
                (a, b, sc, conf, json.dumps(crit, ensure_ascii=False), json.dumps(diff, ensure_ascii=False), json.dumps(miss, ensure_ascii=False), "Catálogo; códigos; aplicações (R&D/Motorcraft)")); n += 1; kept.add((a, b))
        stale = [r[0] for r in c.execute("SELECT id,product_a,product_b FROM equivalence_suggestions WHERE status='pendente'") if (r[1], r[2]) not in kept]
        c.executemany("DELETE FROM equivalence_suggestions WHERE id=?", [(i,) for i in stale])  # só pendentes obsoletas; aprovadas/rejeitadas são preservadas
        c.commit(); return {"candidates": len(cand), "saved": n, "removed_stale": len(stale)}
    finally: c.close()

CRITERIA = ["Código em comum", "Aplicação em comum", "Mesmo eixo", "Mesmo lado", "Lmax igual", "Lmin igual", "Mesmo sistema", "Mesma família"]

def _where(status=None, conf=None, qtxt=None, crit=(), diff=None, minscore=None, skip=(), dif=(), miss=()):
    w, p = [], {}
    if status and "status" not in skip: w.append("e.status=:s"); p["s"] = status
    if conf: w.append("e.confidence=:c"); p["c"] = conf
    if qtxt: w.append("(pa.perfect_code LIKE :q OR pb.perfect_code LIKE :q)"); p["q"] = f"%{qtxt}%"
    if diff == "sem": w.append("e.differences='[]'")
    elif diff == "com": w.append("e.differences<>'[]'")
    if minscore not in (None, ""): w.append("e.score>=:ms"); p["ms"] = float(minscore)
    if "crit" not in skip:
        for i, c in enumerate(x for x in crit if x in CRITERIA): w.append(f"instr(e.criteria,:c{i})>0"); p[f"c{i}"] = '"' + c
    if "dif" not in skip:
        for i, c in enumerate(x for x in dif if x in DIFFS): w.append(f"instr(e.differences,:d{i})>0"); p[f"d{i}"] = '"' + c
    if "miss" not in skip:
        for i, c in enumerate(x for x in miss if x in MISSING): w.append(f"instr(e.missing,:m{i})>0"); p[f"m{i}"] = '"' + c
    return " AND ".join(w) or "1=1", p

DIFFS = ["Nenhuma aplicação em comum", "Eixo diferente", "Lados diferentes", "São lados opostos", "Lmax difere", "Lmin difere", "Sistema diferente", "Família diferente"]
MISSING = ["Eixo", "Lado", "Lmax", "Lmin", "Sistema óleo/gás", "Aplicações", "Diâmetro da haste"]
_BASE = "FROM equivalence_suggestions e JOIN products pa ON pa.id=e.product_a JOIN products pb ON pb.id=e.product_b WHERE "

def listing(status=None, conf=None, qtxt=None, crit=(), diff=None, minscore=None, page=1, size=40, all_rows=False, dif=(), miss=()):
    args = (status, conf, qtxt, crit, diff, minscore); kw = dict(dif=dif, miss=miss); sw, p = _where(*args, **kw); c = db.conn()
    try:
        total = c.execute("SELECT COUNT(*) " + _BASE + sw, p).fetchone()[0]
        lim = "" if all_rows else f"LIMIT {size} OFFSET {(page - 1) * size}"
        rows = [dict(r) for r in c.execute(f"SELECT e.*,pa.perfect_code a_code,pa.description a_desc,pb.perfect_code b_code,pb.description b_desc {_BASE}{sw} ORDER BY e.score DESC,e.id {lim}", p)]
        facets = status_counts = None
        if not all_rows:  # contagens ao lado de cada filtro (respeitando os demais filtros)
            facets = {}
            for grp, col, lst in (("crit", "criteria", CRITERIA), ("dif", "differences", DIFFS), ("miss", "missing", MISSING)):
                w2, p2 = _where(*args, skip=(grp,), **kw)
                for k in lst: facets[grp + ":" + k] = c.execute("SELECT COUNT(*) " + _BASE + w2 + f" AND instr(e.{col},:k)>0", {**p2, "k": '"' + k}).fetchone()[0]
            w3, p3 = _where(*args, skip=("status",), **kw); status_counts = {r[0]: r[1] for r in c.execute("SELECT e.status,COUNT(*) " + _BASE + w3 + " GROUP BY e.status", p3)}
    finally: c.close()
    for r in rows:
        for k in ("criteria", "differences", "missing"): r[k] = json.loads(r[k] or "[]")
    return {"total": total, "rows": rows, "size": size, "facets": facets, "status_counts": status_counts}

def export_xlsx(**f):
    """Planilha com TODAS as sugestões do filtro atual (não só a página). Códigos como texto."""
    import io, datetime
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    rows = listing(all_rows=True, **f)["rows"]; wb = Workbook(); ws = wb.active; ws.title = "Equivalências"
    head = ["Produto A", "Descrição A", "Produto B", "Descrição B", "Similaridade (%)", "Confiança", "Status", "Critérios atendidos", "Diferenças", "Dados ausentes", "Fontes"]
    ws.append(head)
    for r in rows: ws.append([r["a_code"], r["a_desc"], r["b_code"], r["b_desc"], r["score"], r["confidence"], r["status"], "; ".join(r["criteria"]), "; ".join(r["differences"]), "; ".join(r["missing"]), r["sources"]])
    for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="243242"); c.alignment = Alignment(vertical="center", wrap_text=True)
    for i, w in enumerate([14, 36, 14, 36, 12, 11, 11, 60, 50, 50, 40], 1): ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row[:4] + row[7:]: c.number_format = "@"
        for c in row[7:10]: c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    n = wb.create_sheet("Filtros aplicados")
    for k, v in [("Gerado em", datetime.datetime.now().strftime("%d/%m/%Y %H:%M")), ("Registros", len(rows)), ("Status", f.get("status") or "todos"), ("Confiança", f.get("conf") or "todas"),
                 ("Código PERFECT", f.get("qtxt") or "—"), ("Critérios (todos)", ", ".join(f.get("crit") or []) or "—"), ("Diferenças exigidas", ", ".join(f.get("dif") or []) or "—"), ("Dados ausentes exigidos", ", ".join(f.get("miss") or []) or "—"), ("Diferenças", {"sem": "sem diferenças", "com": "com diferenças"}.get(f.get("diff"), "—")),
                 ("Pontuação mínima", f.get("minscore") or "—"), ("Aviso", "A pontuação mede similaridade dos dados e NÃO comprova intercambialidade. Nenhuma sugestão é aprovada automaticamente.")]: n.append([k, v])
    n.column_dimensions["A"].width = 20; n.column_dimensions["B"].width = 90
    for c in n["A"]: c.font = Font(bold=True)
    bio = io.BytesIO(); wb.save(bio); bio.seek(0); return bio, len(rows)
